from pathlib import Path

from docx import Document


path = Path(r"C:\Users\hau28\Desktop\BTL-MNM.docx")
document = Document(str(path))
print("PARAGRAPHS", len(document.paragraphs), "TABLES", len(document.tables), "SECTIONS", len(document.sections))
for index, paragraph in enumerate(document.paragraphs):
    text = paragraph.text.strip()
    if text:
        print(f"P{index:03d} [{paragraph.style.name}] {text}")
for table_index, table in enumerate(document.tables):
    print(f"TABLE {table_index} rows={len(table.rows)} cols={len(table.columns)}")
    for row in table.rows[:8]:
        print(" | ".join(cell.text.replace("\n", " / ").strip() for cell in row.cells))
    if len(table.rows) > 8:
        print("...")
