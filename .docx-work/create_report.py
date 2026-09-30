from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"E:\Learn JavaSpringBoot with ReactJs\Billing-app")
OUT = ROOT / "Bao-cao-BTL-POS-Quan-ly-ban-hang.docx"


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_borders(cell, color="D9D9D9", size="6"):
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn("w:" + name))
        if node is None:
            node = OxmlElement("w:" + name)
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph):
    properties = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    properties.append(keep)


def add_field(paragraph, instruction, display=""):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = display
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Trang ")
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)
    add_field(paragraph, "PAGE", "1")


def add_toc(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Mở tài liệu bằng Word và nhấn Ctrl+A, F9 để cập nhật mục lục."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, placeholder, end])


def add_caption(doc, text, kind="Hình"):
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def add_para(doc, text, bold_lead=None):
    p = doc.add_paragraph(style="Normal")
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_bullets(doc, items, numbered=False):
    style = "List Number" if numbered else "List Bullet"
    for item in items:
        p = doc.add_paragraph(style=style)
        p.add_run(item)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, value in enumerate(headers):
        cell = header.cells[idx]
        cell.text = str(value)
        set_cell_shading(cell, "1F4E78")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(10.5)
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_borders(cell)
    for row_idx, row_values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            cell = cells[idx]
            cell.text = str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_idx % 2 == 1:
                set_cell_shading(cell, "F2F6FA")
            set_cell_borders(cell)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(10.2)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_code_block(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.7)
    p.paragraph_format.right_indent = Cm(0.7)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    for index, line in enumerate(text.splitlines()):
        if index:
            p.add_run().add_break()
        run = p.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(45, 45, 45)
    return p


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(13)
    normal.paragraph_format.line_spacing = 1.3
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name, size, color in (("Title", 22, "000000"), ("Heading 1", 16, "000000"), ("Heading 2", 14, "000000"), ("Heading 3", 13, "000000")):
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
    styles["Caption"].font.name = "Times New Roman"
    styles["Caption"].font.size = Pt(11)
    styles["Caption"].font.italic = True


def add_header_footer(doc):
    for section in doc.sections:
        section.different_first_page_header_footer = True
        header = section.header
        p = header.paragraphs[0]
        p.text = "BÁO CÁO BÀI TẬP LỚN | HỆ THỐNG POS QUẢN LÝ BÁN HÀNG"
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(90, 90, 90)
        footer = section.footer
        fp = footer.paragraphs[0]
        add_page_number(fp)


def page_break(doc):
    doc.add_page_break()


doc = Document()
configure_styles(doc)
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.left_margin = Cm(3.0)
section.right_margin = Cm(2.2)
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(2.0)
section.header_distance = Cm(0.8)
section.footer_distance = Cm(1.0)
add_header_footer(doc)

# Cover page
for _ in range(3):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("TRƯỜNG ĐẠI HỌC ................................................").bold = True
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("KHOA / BỘ MÔN ................................................").bold = True
for _ in range(3):
    doc.add_paragraph()
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("BÁO CÁO BÀI TẬP LỚN")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("XÂY DỰNG WEBSITE POS QUẢN LÝ BÁN HÀNG")
r.bold = True
r.font.size = Pt(18)
for _ in range(3):
    doc.add_paragraph()
cover_rows = [
    ("Môn học", "Chưa cung cấp"),
    ("Giảng viên", "Chưa cung cấp"),
    ("Sinh viên", "Chưa cung cấp"),
    ("Mã sinh viên", "Chưa cung cấp"),
    ("Lớp / Nhóm", "Chưa cung cấp"),
]
t = doc.add_table(rows=0, cols=2)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
for label, value in cover_rows:
    cells = t.add_row().cells
    cells[0].text = label
    cells[1].text = value
    for cell in cells:
        set_cell_borders(cell, color="FFFFFF", size="0")
        set_cell_margins(cell, top=80, bottom=80)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(12)
        cells[0].paragraphs[0].runs[0].bold = True
for _ in range(4):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Thành phố Hồ Chí Minh, năm 2026").italic = True
page_break(doc)

# Opening pages
add_heading(doc, "MỞ ĐẦU", 1)
add_para(doc, "Trong hoạt động bán lẻ, việc quản lý sản phẩm, biến thể, tồn kho, đơn hàng và thanh toán cần được thực hiện nhất quán giữa quầy bán hàng và hệ thống dữ liệu. Đề tài xây dựng website POS quản lý bán hàng nhằm mô hình hóa các nghiệp vụ cốt lõi của một cửa hàng, đồng thời áp dụng kiến thức về phân tích thiết kế hệ thống, lập trình web, cơ sở dữ liệu và kiểm thử.")
add_para(doc, "Báo cáo trình bày kết quả khảo sát và phân tích dự án Billing-app hiện có. Nội dung được xây dựng từ cấu trúc source code, cấu hình, các quyết định kiến trúc và bộ kiểm thử trong repository. Vì vậy, các chức năng được mô tả theo mức độ đã quan sát được trong mã nguồn; những phần mới dừng ở quyết định thiết kế hoặc chưa kiểm chứng runtime được nêu rõ trong phần hạn chế.")
add_heading(doc, "Mục tiêu đề tài", 2)
add_bullets(doc, [
    "Xây dựng giao diện POS hỗ trợ nhân viên thực hiện thao tác bán hàng và theo dõi đơn hàng.",
    "Quản lý người dùng, danh mục, sản phẩm, biến thể, modifier, khách hàng và khuyến mãi.",
    "Tổ chức tồn kho theo sổ cái giao dịch để có thể truy vết nhập, xuất và điều chỉnh.",
    "Tích hợp thanh toán PayOS và hỗ trợ chuyển đổi đơn PENDING sang thanh toán tiền mặt.",
    "Áp dụng JWT, phân quyền admin/staff và activity log cho các thao tác nghiệp vụ quan trọng.",
])
add_heading(doc, "Phạm vi báo cáo", 2)
add_para(doc, "Báo cáo tập trung vào kiến trúc, nghiệp vụ, mô hình dữ liệu, các luồng xử lý chính, giao diện và kết quả kiểm thử của hai thành phần Front-end và billingsoftware. Thông tin về triển khai production, số liệu vận hành thực tế, danh tính nhóm và phân công cá nhân không có trong source được để trống hoặc ghi là chưa cung cấp.")
page_break(doc)

add_heading(doc, "MỤC LỤC", 1)
add_toc(doc.add_paragraph())
add_para(doc, "Ghi chú: khi mở tài liệu bằng Microsoft Word, chọn toàn bộ tài liệu và nhấn Ctrl+A, F9 để cập nhật số trang của mục lục.")
page_break(doc)

add_heading(doc, "CÁC TỪ VIẾT TẮT", 1)
add_table(doc, ["STT", "Viết tắt", "Giải thích"], [
    ("1", "API", "Application Programming Interface - giao diện lập trình ứng dụng"),
    ("2", "CRUD", "Create, Read, Update, Delete - các thao tác dữ liệu cơ bản"),
    ("3", "DTO", "Data Transfer Object - đối tượng trao đổi dữ liệu"),
    ("4", "HTTP", "Hypertext Transfer Protocol"),
    ("5", "JWT", "JSON Web Token"),
    ("6", "ORM", "Object Relational Mapping"),
    ("7", "POS", "Point of Sale - điểm bán hàng"),
    ("8", "PWA", "Progressive Web App"),
    ("9", "SKU", "Stock Keeping Unit - mã quản lý biến thể hàng hóa"),
    ("10", "UI", "User Interface - giao diện người dùng"),
], [1.2, 2.4, 11.5])
page_break(doc)

add_heading(doc, "DANH MỤC HÌNH ẢNH", 1)
add_bullets(doc, [
    "Hình 1. Kiến trúc logic của hệ thống Billing-app",
    "Hình 2. Luồng xác thực và làm mới access token",
    "Hình 3. Luồng tạo đơn hàng và thanh toán",
    "Hình 4. Luồng bù trừ khi hủy hoặc xóa đơn PENDING",
    "Hình 5. Mô hình dữ liệu khái quát",
])
page_break(doc)

add_heading(doc, "DANH MỤC BẢNG BIỂU", 1)
add_bullets(doc, [
    "Bảng 1. Công nghệ sử dụng",
    "Bảng 2. Nhóm người dùng và quyền hạn",
    "Bảng 3. Các nhóm chức năng nghiệp vụ",
    "Bảng 4. Các thực thể dữ liệu chính",
    "Bảng 5. Các endpoint tiêu biểu",
    "Bảng 6. Kết quả kiểm thử và build",
    "Bảng 7. Rủi ro và hướng khắc phục",
])
page_break(doc)

add_heading(doc, "BẢNG PHÂN CÔNG NHIỆM VỤ", 1)
add_para(doc, "Thông tin thành viên và phân công cá nhân chưa được cung cấp trong repository. Bảng dưới đây được giữ làm mẫu để nhóm bổ sung trước khi nộp báo cáo.")
add_table(doc, ["STT", "Họ và tên", "Mã sinh viên", "Nhiệm vụ", "Mức độ hoàn thành"], [
    ("1", "Chưa cung cấp", "Chưa cung cấp", "Phân tích và thiết kế", "Chưa xác định"),
    ("2", "Chưa cung cấp", "Chưa cung cấp", "Lập trình frontend", "Chưa xác định"),
    ("3", "Chưa cung cấp", "Chưa cung cấp", "Lập trình backend và CSDL", "Chưa xác định"),
    ("4", "Chưa cung cấp", "Chưa cung cấp", "Kiểm thử và hoàn thiện báo cáo", "Chưa xác định"),
], [1.0, 3.4, 2.8, 5.0, 3.0])
page_break(doc)

# Chapter 1
add_heading(doc, "CHƯƠNG 1. TÌM HIỂU CÁC CÔNG NGHỆ SỬ DỤNG", 1)
add_heading(doc, "1.1. Tổng quan hệ thống", 2)
add_para(doc, "Billing-app là một ứng dụng web quản lý bán hàng được tổ chức theo mô hình frontend - backend. Front-end cung cấp giao diện React/Vite cho người dùng cuối và nhân viên quản trị. billingsoftware cung cấp REST API bằng Java Spring Boot, thực hiện xác thực, xử lý nghiệp vụ, truy cập MySQL và tích hợp với AWS S3 và PayOS.")
add_para(doc, "Việc tách hai phần giúp giao diện và API có thể phát triển độc lập theo hợp đồng JSON. Frontend không dùng thư viện chia sẻ kiểu tự động với backend; các service Axios ánh xạ dữ liệu DTO từ API thành dữ liệu sử dụng trong component và custom hook.")
add_heading(doc, "1.2. Công nghệ sử dụng", 2)
add_table(doc, ["Thành phần", "Công nghệ", "Vai trò trong hệ thống"], [
    ("Frontend", "React 19, Vite", "Xây dựng giao diện và đóng gói ứng dụng web"),
    ("Điều hướng", "React Router", "Định nghĩa route, route bảo vệ và route admin"),
    ("Server state", "TanStack React Query v5", "Fetch, cache, retry và đồng bộ dữ liệu API"),
    ("HTTP", "Axios", "Gọi REST API, gắn Bearer token và refresh session"),
    ("Backend", "Java 25, Spring Boot 4.1", "Xây dựng REST API và lớp nghiệp vụ"),
    ("Persistence", "Spring Data JPA, Hibernate", "Ánh xạ entity và truy vấn MySQL"),
    ("Bảo mật", "Spring Security, JWT, BCrypt", "Xác thực, phân quyền và mã hóa mật khẩu"),
    ("Thanh toán", "PayOS Java SDK", "Tạo, hủy và xác thực payment link/webhook"),
    ("Lưu trữ", "AWS S3 SDK", "Lưu trữ tệp/hình ảnh sản phẩm"),
], [3.2, 4.3, 7.0])
add_caption(doc, "Bảng 1. Công nghệ sử dụng")
add_heading(doc, "1.3. Kiến trúc logic", 2)
add_code_block(doc, "Người dùng\n    |\n    v\nReact/Vite Front-end\n    | Axios + JSON + Bearer access token\n    v\nSpring Boot REST API /api/v1.0\n    | Security + Service + JPA\n    v\nMySQL\n\nTích hợp ngoài: AWS S3 và PayOS")
add_caption(doc, "Hình 1. Kiến trúc logic của hệ thống Billing-app")
add_para(doc, "Ở phía backend, controller tiếp nhận request và ủy quyền cho service. Service phối hợp repository, entity và các tích hợp ngoài. Ở phía frontend, pages tổ chức theo màn hình; features chứa component và hook theo nghiệp vụ; services tập trung các API call; utils chứa cấu hình Axios, React Query và quản lý session.")
add_heading(doc, "1.4. Công nghệ frontend", 2)
add_para(doc, "React được dùng để xây dựng component và layout. Vite đảm nhiệm dev server và production build. Các trang quản trị được lazy load trong App.jsx để giảm phần code tải ban đầu. ProtectedRoute kiểm tra session, còn AdminRoute kiểm tra vai trò trước khi cho phép truy cập nhóm màn hình quản trị.")
add_para(doc, "TanStack React Query được dùng cho server state. Cấu hình hiện tại đặt stale time mặc định một giờ, không retry với lỗi 401/403 và retry tối đa ba lần với các lỗi khác. Đây là cách tách server state khỏi state giao diện và tránh tự viết cơ chế cache riêng cho từng component.")
add_heading(doc, "1.5. Công nghệ backend", 2)
add_para(doc, "Spring Boot cung cấp nền tảng tạo REST API, Spring Data JPA cung cấp repository và Spring Security xử lý chuỗi bảo mật. Backend đặt context path /api/v1.0, nhận thông tin kết nối MySQL và khóa tích hợp từ biến môi trường. JPA hiện dùng ddl-auto=update; đây là lựa chọn thuận tiện cho phát triển nhưng cần thay bằng công cụ migration có phiên bản khi triển khai production.")
page_break(doc)

# Chapter 2
add_heading(doc, "CHƯƠNG 2. TÌM HIỂU BÀI TOÁN VÀ PHÂN TÍCH THIẾT KẾ HỆ THỐNG", 1)
add_heading(doc, "2.1. Mục đích và phạm vi", 2)
add_para(doc, "Bài toán đặt ra là xây dựng hệ thống hỗ trợ cửa hàng quản lý hàng hóa và bán hàng tại điểm bán. Hệ thống phải cho phép nhân viên đăng nhập, chọn sản phẩm, tạo order, nhận thanh toán và theo dõi trạng thái. Người quản trị cần quản lý dữ liệu nền, tồn kho, khuyến mãi, khách hàng và lịch sử thao tác.")
add_para(doc, "Phạm vi triển khai thực tế trong repository gồm quản lý user, category, item, variant, modifier group, promotion, customer, order, inventory, dashboard và activity log. Hệ thống có tích hợp PayOS, nhưng các thông tin môi trường triển khai thật và dữ liệu vận hành không thuộc phạm vi source code khảo sát.")
add_heading(doc, "2.2. Nhóm người dùng và quyền hạn", 2)
add_table(doc, ["Nhóm", "Quyền chính", "Màn hình/endpoint tiêu biểu"], [
    ("Nhân viên", "Đăng nhập, bán hàng, xem đơn, xem dữ liệu nghiệp vụ được cấp quyền", "Dashboard, Explore, Orders, Activity logs"),
    ("Quản trị viên", "Toàn bộ quyền nhân viên và CRUD dữ liệu quản trị", "Items, Categories, Modifiers, Users, Promotions, Customers"),
    ("Hệ thống PayOS", "Gửi webhook kết quả thanh toán", "POST /payos/webhook"),
], [3.0, 6.1, 5.4])
add_caption(doc, "Bảng 2. Nhóm người dùng và quyền hạn")
add_heading(doc, "2.3. Phân tích chức năng", 2)
add_table(doc, ["Nhóm chức năng", "Chức năng cụ thể"], [
    ("Xác thực", "Login, logout, refresh access token, mã hóa mật khẩu, bảo vệ route"),
    ("Danh mục và hàng hóa", "CRUD category, item, variant, ảnh và modifier group"),
    ("Bán hàng", "Tìm hàng, chọn biến thể/modifier, gắn khách hàng, áp mã coupon, tạo order"),
    ("Thanh toán", "CASH, PayOS QR, webhook hoàn tất, hủy payment link, switch-to-cash"),
    ("Tồn kho", "Nhập/xuất/điều chỉnh, stock check, xem lịch sử ledger, cập nhật cache"),
    ("Khuyến mãi", "Tạo, sửa, bật/tắt, xóa, đánh giá promotion tốt nhất và giới hạn sử dụng"),
    ("Báo cáo", "Dashboard metric, order history, phân trang và lọc"),
    ("Audit", "Ghi nhận login và thay đổi dữ liệu nghiệp vụ, lọc theo vai trò/ngày/action"),
], [4.0, 10.5])
add_caption(doc, "Bảng 3. Các nhóm chức năng nghiệp vụ")
add_heading(doc, "2.4. Thiết kế luồng xác thực", 2)
add_code_block(doc, "Login -> kiểm tra email/mật khẩu -> tạo JWT access token\n      -> issue refresh token -> lưu hash -> gửi cookie HttpOnly\nRequest -> Axios gắn Bearer token -> JwtRequestFilter xác thực\n401 đủ điều kiện -> gọi /auth/refresh -> rotate refresh token\nRefresh lỗi -> clear session + clear query cache -> /login")
add_caption(doc, "Hình 2. Luồng xác thực và làm mới access token")
add_para(doc, "Access token được giữ trong memory ở frontend. Refresh token được backend xử lý qua cookie HttpOnly và có cơ chế rotate/revoke. Khi request hết hạn, interceptor dùng refresh promise dùng chung để tránh nhiều request đồng thời cùng tạo nhiều phiên refresh.")
add_heading(doc, "2.5. Thiết kế luồng tạo đơn và thanh toán", 2)
add_code_block(doc, "Cart -> server evaluate promotion và tính lại total\n     -> liên kết Customer và cập nhật CRM\n     -> tăng timesUsed của Promotion nếu hợp lệ\n     -> lưu Order và OrderItem\n     -> ghi InventoryTransaction OUT, cập nhật cached stock\n     -> CASH: COMPLETED\n     -> PAYOS: tạo payment link, chờ webhook COMPLETED")
add_caption(doc, "Hình 3. Luồng tạo đơn hàng và thanh toán")
add_para(doc, "Việc tính lại tổng tiền ở server là điểm quan trọng để không phụ thuộc hoàn toàn vào dữ liệu giá do trình duyệt gửi. Order service cũng ghi nhận giao dịch tồn kho và log tạo đơn trong cùng quy trình nghiệp vụ.")
add_heading(doc, "2.6. Thiết kế luồng bù trừ", 2)
add_code_block(doc, "PENDING + cancel/delete\n    -> tạo InventoryTransaction IN\n    -> tính lại cachedStockQuantity\n    -> hoàn timesUsed promotion\n    -> hoàn orderCount/totalSpent của customer\n    -> hủy payment link PayOS\n    -> CANCELLED hoặc xóa order\nCOMPLETED -> không cho xóa")
add_caption(doc, "Hình 4. Luồng bù trừ khi hủy hoặc xóa đơn PENDING")
add_heading(doc, "2.7. Thiết kế dữ liệu", 2)
add_code_block(doc, "Category 1---N Item 1---N Variant 1---N InventoryTransaction\n                         |\n                         +---N ModifierGroup --- N Modifier\nCustomer 1---N Order 1---N OrderItem N---1 Variant/Item\nPromotion 1---N Order (tham chiếu promotionId)\nUser 1---N ActivityLog\nOrder --- PaymentDetails --- PayOS")
add_caption(doc, "Hình 5. Mô hình dữ liệu khái quát")
add_table(doc, ["Thực thể", "Thuộc tính/quan hệ chính", "Mục đích"], [
    ("User", "email, password, role, trạng thái", "Xác thực và phân quyền"),
    ("Category", "categoryId, name, description", "Phân nhóm mặt hàng"),
    ("Item", "itemId, name, price, category", "Sản phẩm cha"),
    ("Variant", "variantId, SKU, basePrice, attributes JSON, cachedStock", "Đơn vị hàng hóa quản lý tồn kho"),
    ("InventoryTransaction", "type, quantity, referenceId, note", "Sổ cái biến động tồn kho"),
    ("Order/OrderItem", "customer, total, payment, danh sách item", "Hóa đơn và chi tiết bán hàng"),
    ("Promotion", "type, usageLimit, timesUsed, active", "Quản lý ưu đãi"),
    ("ActivityLog", "email, action, entity, timestamp", "Audit trail"),
], [3.5, 6.0, 5.0])
add_caption(doc, "Bảng 4. Các thực thể dữ liệu chính")
page_break(doc)

# Chapter 3
add_heading(doc, "CHƯƠNG 3. CÀI ĐẶT CHƯƠNG TRÌNH", 1)
add_heading(doc, "3.1. Cài đặt cơ sở dữ liệu", 2)
add_para(doc, "Backend sử dụng MySQL thông qua Spring Data JPA. Các entity được ánh xạ vào bảng với tiền tố tbl ở một số thực thể quan trọng như tbl_orders, tbl_variants, tbl_inventory_transactions và tbl_activity_logs. Hibernate được cấu hình ddl-auto=update nên schema có thể được cập nhật theo entity trong môi trường phát triển.")
add_para(doc, "Trong một hệ thống thực tế, schema production nên được quản lý bằng migration có version để kiểm soát thay đổi, rollback và lịch sử triển khai. Đây là một hạn chế đã được ghi nhận trong phân tích, không phải chức năng đang có của repository.")
add_heading(doc, "3.2. Cài đặt frontend và các màn hình", 2)
add_para(doc, "App.jsx định nghĩa route login, dashboard, explore, order history, activity logs và các màn hình quản trị. Các trang quản trị được tải lười bằng React.lazy. Layout chung chứa menubar và nội dung bên trong route bảo vệ.")
add_table(doc, ["Màn hình", "Nội dung chính", "Vai trò"], [
    ("Login", "Nhập email/mật khẩu, hiển thị lỗi xác thực", "Nhân viên, quản trị viên"),
    ("Dashboard", "Các metric và đơn gần đây", "Nhân viên, quản trị viên"),
    ("Explore", "Danh mục, item, giỏ hàng, khách hàng, coupon, thanh toán", "Nhân viên, quản trị viên"),
    ("Order History", "Danh sách, tìm kiếm, lọc, xem/hủy/xóa đơn", "Nhân viên, quản trị viên"),
    ("Manage Items/Categories", "CRUD hàng hóa và danh mục", "Quản trị viên"),
    ("Manage Modifiers/Promotions", "CRUD modifier và khuyến mãi", "Quản trị viên"),
    ("Manage Users/Customers", "Quản lý tài khoản và khách hàng", "Quản trị viên"),
    ("Activity Logs", "Xem audit log theo bộ lọc", "Nhân viên, quản trị viên"),
], [4.1, 7.2, 3.2])
add_heading(doc, "3.3. Các service và endpoint tiêu biểu", 2)
add_table(doc, ["Nhóm", "Endpoint tiêu biểu", "Chức năng"], [
    ("Auth", "POST /login; POST /auth/refresh; POST /auth/logout", "Đăng nhập và quản lý phiên"),
    ("Order", "GET/POST /orders; POST /orders/{id}/cancel", "Tra cứu và tạo/hủy order"),
    ("Payment", "POST /payos/webhook", "Nhận kết quả thanh toán PayOS"),
    ("Inventory", "POST /admin/inventory/transactions; stock-check", "Ghi và kiểm kê tồn kho"),
    ("Promotion", "GET /promotions/active; POST /promotions/evaluate", "Đánh giá ưu đãi"),
    ("Audit", "GET /activity-logs", "Lọc và phân trang log"),
], [3.0, 6.4, 5.1])
add_caption(doc, "Bảng 5. Các endpoint tiêu biểu")
add_heading(doc, "3.4. Cài đặt nghiệp vụ tồn kho", 2)
add_para(doc, "Mỗi biến thể có cachedStockQuantity. Tuy nhiên, giá trị này không phải nguồn lịch sử duy nhất: InventoryTransactionEntity lưu từng lần IN, OUT hoặc ADJUSTMENT với referenceId và note. Khi checkout, hệ thống ghi OUT với số lượng âm; khi hủy/xóa đơn PENDING, hệ thống ghi IN bù trừ. Sau mỗi giao dịch, tồn kho được tính lại và lưu vào cache của variant.")
add_heading(doc, "3.5. Cài đặt khuyến mãi và modifier", 2)
add_para(doc, "Promotion được đánh giá từ cart ở backend. Checkout sử dụng kết quả đánh giá để đặt subtotal, discount, tax, grand total và promotion reference. Modifier được lưu độc lập với giá base của variant, giúp thể hiện lựa chọn bổ sung trong order item mà không biến modifier thành hàng tồn kho riêng.")
add_heading(doc, "3.6. Cài đặt audit log", 2)
add_para(doc, "ActivityLogServiceImpl ghi action, entity type, entity id, mô tả và email người thực hiện. Nếu service không nhận email tường minh, email được lấy từ SecurityContextHolder; khi không có user hợp lệ, code dùng địa chỉ hệ thống. Controller activity log giới hạn phạm vi xem của staff và cho phép admin lọc theo người dùng.")
add_heading(doc, "3.7. Kiểm thử và kết quả thực nghiệm", 2)
add_para(doc, "Repository có bộ test tích hợp backend sử dụng SpringBootTest, MockMvc và Transactional. Các nhóm test bao phủ authentication, CRUD, promotion evaluation, checkout, pagination, inventory ledger, stock check, order compensation, dashboard và activity log. Frontend có test Node.js cho auth message, retry policy, in-memory session, promotion error mapping và variant defaults.")
add_table(doc, ["Hạng mục", "Kết quả khảo sát", "Nhận xét"], [
    ("Frontend build", "Đạt", "npm run build hoàn thành với Vite"),
    ("Frontend test", "11/12 đạt", "Một test auth-session thất bại do implementation có thêm name: null"),
    ("Backend test", "Chưa có kết quả chạy", "mvnw.cmd không khởi chạy được Maven wrapper trong môi trường khảo sát"),
    ("Runtime integration", "Chưa xác nhận", "Cần MySQL và các biến môi trường PayOS/AWS/JWT"),
], [3.5, 3.8, 7.2])
add_caption(doc, "Bảng 6. Kết quả kiểm thử và build")
add_para(doc, "Kết quả trên không đồng nghĩa với việc backend có lỗi nghiệp vụ. Riêng backend chưa thể chạy do giới hạn môi trường công cụ; vì vậy báo cáo không đưa ra kết luận pass/fail cho toàn bộ backend test suite.")
add_heading(doc, "3.8. Rủi ro và hướng khắc phục", 2)
add_table(doc, ["Rủi ro", "Ảnh hưởng", "Hướng khắc phục đề xuất"], [
    ("ddl-auto=update", "Khó kiểm soát schema production", "Dùng Flyway/Liquibase và migration có version"),
    ("URL PayOS còn phụ thuộc môi trường", "Có thể trả về localhost khi triển khai", "Đưa return/cancel URL vào cấu hình theo môi trường"),
    ("Offline-first chưa có bằng chứng đầy đủ", "Không nên cam kết PWA offline hoàn chỉnh", "Bổ sung service worker, IndexedDB, queue và cơ chế sync"),
    ("Một test session bất nhất contract", "CI frontend có thể fail", "Thống nhất schema AuthResponse và test fixture"),
    ("Backend chưa chạy test trong môi trường này", "Chưa có bằng chứng runtime hiện tại", "Cài Maven wrapper/dependency và cấu hình DB test trong CI"),
], [3.5, 5.0, 6.0])
add_caption(doc, "Bảng 7. Rủi ro và hướng khắc phục")
page_break(doc)

# Conclusion and references
add_heading(doc, "KẾT LUẬN", 1)
add_heading(doc, "Kết quả đạt được", 2)
add_para(doc, "Qua quá trình phân tích, dự án đã hình thành một nền tảng POS có các nhóm chức năng tương đối đầy đủ cho bài toán quản lý bán hàng: đăng nhập và phân quyền; quản lý danh mục, mặt hàng, biến thể và modifier; tạo đơn; thanh toán PayOS/CASH; quản lý khách hàng và khuyến mãi; tồn kho theo ledger; dashboard; và activity log.")
add_para(doc, "Về kỹ thuật, dự án thể hiện cách tổ chức frontend React/Vite và backend Spring Boot theo hai thư mục độc lập, có lớp service/repository, có route guard, có cơ chế refresh session, có đánh giá promotion phía server và có quy trình bù trừ để bảo vệ tính nhất quán khi hủy hoặc xóa order PENDING.")
add_heading(doc, "Hạn chế của hệ thống", 2)
add_bullets(doc, [
    "Schema vẫn dùng ddl-auto=update và chưa có migration versioned.",
    "Cấu hình production, URL PayOS, CORS và cookie cần được kiểm soát theo từng môi trường.",
    "ADR về offline-first/IndexedDB chưa được chứng minh đầy đủ bằng implementation frontend hiện tại.",
    "Frontend còn một test không đồng nhất với shape session thực tế.",
    "Backend test suite chưa được chạy thành công trong môi trường khảo sát do Maven wrapper.",
])
add_heading(doc, "Hướng phát triển", 2)
add_bullets(doc, [
    "Bổ sung Flyway hoặc Liquibase, chuẩn hóa quy trình backup và rollback schema.",
    "Hoàn thiện offline queue, IndexedDB, service worker và cơ chế đồng bộ có kiểm soát xung đột.",
    "Chuẩn hóa API response và pagination metadata mà không phá vỡ các client hiện tại.",
    "Bổ sung CI chạy lint, frontend test, backend test và kiểm tra migration.",
    "Tách cấu hình môi trường, tăng cường logging có cấu trúc và giám sát các tích hợp PayOS/AWS.",
])
add_heading(doc, "TÀI LIỆU THAM KHẢO", 1)
refs = [
    "Repository Billing-app, CONTEXT.md và source code Front-end/billingsoftware, truy cập ngày 13/09/2026.",
    "Spring Boot Reference Documentation, https://docs.spring.io/spring-boot/.",
    "Spring Security Reference, https://docs.spring.io/spring-security/reference/.",
    "Spring Data JPA Reference Documentation, https://docs.spring.io/spring-data/jpa/reference/.",
    "React Documentation, https://react.dev/.",
    "Vite Documentation, https://vite.dev/.",
    "TanStack Query Documentation, https://tanstack.com/query/latest.",
    "PayOS Documentation, https://payos.vn/docs/.",
]
for ref in refs:
    p = doc.add_paragraph(style="List Number")
    p.add_run(ref)

# Update fields on open in Word.
settings = doc.settings._element
update = settings.find(qn("w:updateFields"))
if update is None:
    update = OxmlElement("w:updateFields")
    settings.append(update)
update.set(qn("w:val"), "true")

doc.core_properties.title = "Báo cáo bài tập lớn xây dựng website POS quản lý bán hàng"
doc.core_properties.subject = "Phân tích và triển khai Billing-app"
doc.core_properties.author = ""
doc.core_properties.comments = ""
doc.save(str(OUT))
print(OUT)
